"""Create the submission report from measured training results: python make_report.py."""
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches

import report_text as rt
from report_text import draw_architecture_diagram


def main():
    output = Path("results")
    metrics = json.loads((output / "metrics.json").read_text(encoding="utf-8"))
    test = metrics["test"]
    tuning_path = output / "tuning.json"
    tuning = json.loads(tuning_path.read_text(encoding="utf-8")) if tuning_path.exists() else None
    final, final_label = rt.resolve_final(metrics, tuning)
    draw_architecture_diagram(output / "architecture_diagram.png", dropout=metrics.get("dropout", 0.3))
    doc = Document()
    doc.add_heading("Image Classification with Convolutional Neural Networks", 0)
    doc.add_paragraph("Neural Networks Implementation Project Report")

    figure_n = 0

    def add_figure(image_name, caption, width):
        nonlocal figure_n
        figure_n += 1
        doc.add_picture(str(output / image_name), width=Inches(width))
        doc.add_paragraph(f"Figure {figure_n}. {caption}")

    def add_table(headers, rows):
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Shading Accent 1"
        for cell, title in zip(table.rows[0].cells, headers):
            cell.text = title
        for row in rows:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = str(value)

    # --- Introduction: Problem description and dataset overview ---
    doc.add_heading("Introduction", 1)
    doc.add_paragraph(rt.introduction(metrics))
    add_table(("Class", "Images"), list(metrics["class_counts"].items()))

    # --- Methodology ---
    doc.add_heading("Methodology", 1)
    doc.add_heading("Preprocessing steps", 2)
    doc.add_paragraph(rt.preprocessing(metrics))

    doc.add_heading("CNN architecture design", 2)
    doc.add_paragraph(rt.architecture(metrics))
    add_figure("architecture_diagram.png", "Shared CNN architecture. The regularized variant additionally applies batch normalization and dropout as noted in each block.", 7.0)

    doc.add_heading("Training setup and hyperparameters", 2)
    doc.add_paragraph(rt.training_setup(metrics))

    # --- Results ---
    doc.add_heading("Results", 1)
    doc.add_heading("Performance metrics", 2)
    for line in rt.results_summary_lines(metrics, test):
        doc.add_paragraph(line)
    doc.add_paragraph("Macro metrics average the eight classes equally. The test accuracy uses all held-out examples. Per-class precision, recall, and F1-score are shown below.")
    add_table(
        ("Class", "Precision", "Recall", "F1-score", "Support"),
        [(name, f"{test['classification_report'][name]['precision']:.3f}", f"{test['classification_report'][name]['recall']:.3f}",
          f"{test['classification_report'][name]['f1-score']:.3f}", int(test['classification_report'][name]['support']))
         for name in metrics["class_counts"]],
    )

    doc.add_heading("Confusion matrix and classification report", 2)
    add_figure("confusion_matrix.png", "Confusion matrix for the selected model on the held-out test set. Rows are true labels; columns are predictions.", 5.6)

    doc.add_heading("Training/validation loss and accuracy curves", 2)
    add_figure("training_curves.png", "Training and validation loss and accuracy for both CNN variants. Each dashed line is validation performance.", 6.2)

    # --- Discussion ---
    doc.add_heading("Discussion", 1)
    doc.add_heading("Observations, challenges, and insights", 2)
    for paragraph in rt.observations_paragraphs(metrics, final):
        doc.add_paragraph(paragraph)

    doc.add_heading("Justification for design and optimization decisions", 2)
    doc.add_paragraph(rt.justification_intro(metrics))
    if tuning:
        doc.add_paragraph(rt.tuning_intro(tuning))
        add_table(
            ("Learning rate", "Best validation accuracy"),
            [(lr, f"{tuning['learning_rate_sweep'][str(lr)]:.3%}") for lr in tuning["learning_rates_tried"]],
        )
        doc.add_paragraph(f"Best learning rate: {tuning['best_learning_rate']}.")
        add_table(
            ("Dropout rate", "Best validation accuracy"),
            [(rate, f"{tuning['dropout_sweep_at_best_lr'][str(rate)]:.3%}") for rate in tuning["dropout_rates_tried"]],
        )
        doc.add_paragraph(f"Best dropout rate at learning rate {tuning['best_learning_rate']}: {tuning['best_dropout']}.")
        add_figure("tuning_curves.png", "Best validation accuracy across the learning rate sweep (left) and the dropout sweep at the best learning rate (right).", 6.2)
        doc.add_paragraph(rt.tuning_outcome(metrics, tuning, final))
        note = rt.tuning_variance_note(tuning, metrics)
        if note:
            doc.add_paragraph(note)

    doc.add_heading("Potential improvements and future work", 2)
    doc.add_paragraph(rt.future_work())

    # --- Final insight and recommendations ---
    doc.add_heading("Final Insight and Recommendations", 1)
    for paragraph in rt.final_insight_and_recommendations(metrics, tuning, final, final_label):
        doc.add_paragraph(paragraph)

    doc.add_heading("Reproducibility and source", 1)
    doc.add_paragraph(rt.reproducibility(metrics))

    # --- Conclusion: always last ---
    doc.add_heading("Conclusion", 1)
    doc.add_paragraph(rt.conclusion(metrics, tuning, final, final_label))

    doc.save("CNN_Project_Report.docx")
    print("Created CNN_Project_Report.docx")


if __name__ == "__main__":
    main()
